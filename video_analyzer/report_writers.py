"""
Write an analysis report to disk in the user's chosen format:
Markdown (.md), Word (.docx), or PDF (.pdf).

The .docx and .pdf writers import their libraries lazily so the tool still
runs (Markdown always works) even if those extras aren't installed.
"""

import re
from pathlib import Path
from typing import Optional

# format key -> file extension
FORMATS = {
    "md": ".md",
    "docx": ".docx",
    "pdf": ".pdf",
    "txt": ".txt",
    "srt": ".srt",
}

# Friendly labels for UIs
FORMAT_LABELS = {
    "Markdown (.md)": "md",
    "Word document (.docx)": "docx",
    "PDF (.pdf)": "pdf",
}


def _header_md(meta: dict) -> str:
    return (
        f"# Video Analysis: {meta['filename']}\n\n"
        f"**Model**: {meta.get('model', '')}  \n"
        f"**File**: {meta['filename']}  \n"
        f"**Size**: {meta.get('size_mb', 0):.1f} MB  \n\n---\n\n"
    )


# --------------------------------------------------------------------------- md
def write_markdown(out_file: Path, analysis: str, meta: dict) -> None:
    out_file.write_text(_header_md(meta) + analysis, encoding="utf-8")


# ------------------------------------------------------------------------- docx
def _add_runs_with_bold(paragraph, text: str):
    """Add text to a docx paragraph, honoring **bold** markers."""
    for i, chunk in enumerate(text.split("**")):
        if chunk:
            run = paragraph.add_run(chunk)
            run.bold = i % 2 == 1


def _render_md_into_docx(doc, analysis: str) -> None:
    """Render lightweight markdown (headings, tables, bullets, bold) into a docx."""
    lines = analysis.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2), level=level)
            i += 1
            continue

        # Tables: collect consecutive lines starting with '|'
        if stripped.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = []
            for row in block:
                # skip separator rows like |---|---|
                if re.match(r"^\|[\s:\-\|]+\|?$", row):
                    continue
                cells = [c.strip() for c in row.strip("|").split("|")]
                rows.append(cells)
            if rows:
                ncols = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=ncols)
                try:
                    table.style = "Light Grid Accent 1"
                except Exception:
                    pass
                for r in rows:
                    cells = table.add_row().cells
                    for j in range(ncols):
                        cells[j].text = r[j] if j < len(r) else ""
            continue

        # Bullets
        mb = re.match(r"^[-*]\s+(.*)$", stripped)
        if mb:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs_with_bold(p, mb.group(1))
            i += 1
            continue

        # Horizontal rule -> skip
        if re.match(r"^-{3,}$", stripped):
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        _add_runs_with_bold(p, stripped)
        i += 1


def write_docx(out_file: Path, analysis: str, meta: dict) -> None:
    from docx import Document  # lazy import

    doc = Document()
    doc.add_heading(f"Video Analysis: {meta['filename']}", level=0)
    doc.add_paragraph(f"Model: {meta.get('model', '')}")
    doc.add_paragraph(f"File: {meta['filename']}")
    doc.add_paragraph(f"Size: {meta.get('size_mb', 0):.1f} MB")
    _render_md_into_docx(doc, analysis)
    doc.save(str(out_file))


# -------------------------------------------------------------------------- pdf
_PDF_CSS = """
@page { size: letter; margin: 2cm; }
body { font-family: Helvetica, Arial, sans-serif; font-size: 10pt; color: #111; line-height: 1.4; }
h1 { font-size: 18pt; } h2 { font-size: 14pt; } h3 { font-size: 12pt; }
h1, h2, h3 { color: #1a1a2e; }
table { border-collapse: collapse; width: 100%; margin: 8px 0; }
th, td { border: 1px solid #999; padding: 4px 6px; text-align: left; font-size: 9pt; vertical-align: top; }
th { background: #eee; }
hr { border: none; border-top: 1px solid #ccc; }
"""


def write_pdf(out_file: Path, analysis: str, meta: dict) -> None:
    import markdown as _md  # lazy import
    from xhtml2pdf import pisa  # lazy import

    full_md = _header_md(meta) + analysis
    body = _md.markdown(full_md, extensions=["tables", "fenced_code", "sane_lists"])
    html = f"<html><head><meta charset='utf-8'><style>{_PDF_CSS}</style></head><body>{body}</body></html>"
    with open(out_file, "wb") as f:
        result = pisa.CreatePDF(html, dest=f, encoding="utf-8")
    if result.err:
        raise RuntimeError(f"PDF generation failed for {out_file.name}")


# --------------------------------------------------------------------- dispatch
def write_report(out_dir: Path, stem: str, fmt: str, analysis: str, meta: dict) -> Path:
    """Write the report in the given format and return the output path."""
    ext = FORMATS.get(fmt, ".md")
    # transcript exports use a clearer suffix
    suffix = "_transcript" if fmt in ("txt", "srt") else "_analysis"
    out_file = out_dir / f"{stem}{suffix}{ext}"
    if fmt == "docx":
        write_docx(out_file, analysis, meta)
    elif fmt == "pdf":
        write_pdf(out_file, analysis, meta)
    elif fmt in ("txt", "srt"):
        # raw text — no markdown header (SRT must be exactly valid)
        out_file.write_text(analysis, encoding="utf-8")
    else:
        write_markdown(out_file, analysis, meta)
    return out_file


# --------------------------------------------------- compiled (synthesis/hook bank)
def write_compiled(out_file: Path, title: str, body: str) -> None:
    """Write a standalone compiled document (synthesis / hook bank) by extension."""
    ext = out_file.suffix.lower()
    if ext == ".docx":
        from docx import Document
        doc = Document()
        doc.add_heading(title, level=0)
        _render_md_into_docx(doc, body)
        doc.save(str(out_file))
    elif ext == ".pdf":
        import markdown as _md
        from xhtml2pdf import pisa
        html_body = _md.markdown(f"# {title}\n\n{body}",
                                 extensions=["tables", "fenced_code", "sane_lists"])
        html = f"<html><head><meta charset='utf-8'><style>{_PDF_CSS}</style></head><body>{html_body}</body></html>"
        with open(out_file, "wb") as f:
            result = pisa.CreatePDF(html, dest=f, encoding="utf-8")
        if result.err:
            raise RuntimeError(f"PDF generation failed for {out_file.name}")
    else:
        out_file.write_text(f"# {title}\n\n{body}", encoding="utf-8")
